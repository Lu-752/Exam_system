# -*- coding: utf-8 -*-
"""postprocess 增补：表格整块不跨页、单元格行不拆断、标题不与后文分离。"""
import io

P = r"C:\Users\lxq\Desktop\exam-sys\output\20261027-deploy\working\postprocess_toc.py"
s = io.open(P, encoding="utf-8").read()

FUNC = '''def keep_blocks_together(doc):
    """表格整块留在同一页、行不跨页拆断、标题紧跟后文。

    html-to-docx 不做 keepNext，长表会从行中间断开，出现"下一页只剩半张表"。
    Word 的做法是：每行 cantSplit；除最后一行外，所有段落 keepNext。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for tbl in doc.tables:
        rows = tbl.rows
        for idx, row in enumerate(rows):
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(OxmlElement("w:cantSplit"))
            if idx < len(rows) - 1:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True

    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True


'''

s = s.replace("def main():", FUNC + "def main():", 1)
s = s.replace("    apply_breaks(doc)\n    fix_pagenum(doc)\n",
              "    apply_breaks(doc)\n    keep_blocks_together(doc)\n    fix_pagenum(doc)\n", 1)

io.open(P, "w", encoding="utf-8", newline="").write(s)
print("postprocess patched", "keep_blocks_together(doc)" in s)
