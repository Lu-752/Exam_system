# -*- coding: utf-8 -*-
"""考试系统操作演示 PPT 生成器（白底编辑风 / 纯 python-pptx）。

用法：
    python build_deck.py            # 生成 pptx
设计规范见 docs/考试系统演示材料/DESIGN.md
"""
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

BASE = r"C:\Users\lxq\Desktop\exam-sys\docs\考试系统演示材料"
ASSET = os.path.join(BASE, "assets2")
OUT = os.path.join(BASE, "考试系统演示材料.pptx")

W, H, M = 1280, 720, 80
CW = W - 2 * M
PX = 9525  # 1px @96dpi

INK, BODY, MUTED, FAINT = "#17191E", "#4A4F5A", "#8C919C", "#B0B4BD"
HAIR, HAIR2 = "#DCDEE3", "#EDEFF2"
DEEP = "#1F3A5F"
FONT = "Microsoft YaHei"


def px(v):
    return Emu(int(round(v * PX)))


def _set_ea(run, name):
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import nsdecls

        ea = parse_xml("<a:ea %s/>" % nsdecls("a"))
        rpr.append(ea)
    ea.set("typeface", name)


def text(sl, x, y, w, h, s, size, color=INK, bold=False, align="l",
         lh=1.4, anchor="t", space_before=0):
    tb = sl.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT,
                   "c": PP_ALIGN.CENTER}[align]
    p.line_spacing = lh
    if space_before:
        p.space_before = px(space_before)
    r = p.add_run()
    r.text = s
    r.font.size = Pt(size * 0.75)
    r.font.bold = bold
    r.font.name = FONT
    r.font.color.rgb = RGBColor.from_string(color[1:])
    _set_ea(r, FONT)
    return tb


def rich(sl, x, y, w, h, parts, lh=1.4, align="l", space_before=0):
    """parts = [(text, size, color, bold), ...] 同段落多 run"""
    tb = sl.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT,
                   "c": PP_ALIGN.CENTER}[align]
    p.line_spacing = lh
    if space_before:
        p.space_before = px(space_before)
    for s, size, color, bold in parts:
        r = p.add_run()
        r.text = s
        r.font.size = Pt(size * 0.75)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = RGBColor.from_string(color[1:])
        _set_ea(r, FONT)
    return tb


def rule(sl, x, y, w, h=1, color=HAIR):
    from pptx.enum.shapes import MSO_SHAPE

    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color[1:])
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def picture(sl, name, x, y, w, border=True):
    path = os.path.join(ASSET, name)
    iw, ih = Image.open(path).size
    ratio = iw / ih
    h = w / ratio
    pic = sl.shapes.add_picture(path, px(x), px(y), px(w), px(h))
    if border:
        pic.line.color.rgb = RGBColor.from_string(HAIR[1:])
        pic.line.width = Pt(0.75)
    return w, h


def ratio_of(name):
    iw, ih = Image.open(os.path.join(ASSET, name)).size
    return iw / ih


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def footer(sl, label, no, total):
    text(sl, M, 668, 400, 20, label, 11, FAINT)
    text(sl, W - M - 120, 668, 120, 20, "%02d / %d" % (no, total), 11, FAINT, align="r")


def eyebrow(sl, no, tag, y=58):
    rich(sl, M, y, 600, 20, [(no, 12, DEEP, True), ("　" + tag, 12, MUTED, False)])


# ---------------------------------------------------------------- 版式

def slide_cover(prs):
    sl = blank(prs)
    rule(sl, 0, 0, 12, H, DEEP)
    text(sl, M, 96, 700, 22, "产品操作演示", 13, MUTED)
    text(sl, M, 140, 900, 84, "考试系统", 56, INK, True)
    text(sl, M, 240, 860, 30, "十三个动作，走完一场考试", 20, BODY)
    rule(sl, M, 322, 380, 2, DEEP)
    text(sl, M, 350, 460, 24, "管理员：命题、组卷、排考、监考、阅卷", 13, MUTED, lh=1.9)
    text(sl, M, 380, 460, 24, "学员：报名、答题、交卷、查分", 13, MUTED, lh=1.9)
    text(sl, M, 560, 500, 22, "V0.1.0　·　2026 年 10 月", 13, MUTED)
    return sl


def slide_index(prs, items, total):
    sl = blank(prs)
    eyebrow(sl, "目录", "操作步骤")
    text(sl, M, 88, 900, 50, "从建题库到出成绩", 32, INK, True)
    text(sl, M, 146, 900, 24, "01–11 是管理员视角，12–13 是学员视角；全程一份数据，不需要导表。", 13, BODY)

    n = len(items)
    half = 7
    cols = [(M, 520, items[:half]), (M + 560, 520, items[half:])]
    top = 205
    step = 52
    for cx, cw, group in cols:
        for i, (no, name, page) in enumerate(group):
            y = top + i * step
            rich(sl, cx, y, 40, 22, [(no, 12, DEEP, True)])
            text(sl, cx + 44, y - 2, cw - 130, 24, name, 14, INK)
            text(sl, cx + cw - 62, y - 1, 62, 22, page, 12, MUTED, align="r")
            rule(sl, cx, y + 32, cw, 1, HAIR2)
    rule(sl, M, 596, CW, 1, HAIR2)
    text(sl, M, 612, CW, 24, "讲的时候按 01 → 13 的顺序走，每页三步，讲到哪一步就在系统里点哪一步。", 12, MUTED)
    footer(sl, "目录", 2, total)
    return sl


def slide_split(prs, no, tag, title, path, steps, img, side, note, label,
                page, total):
    sl = blank(prs)
    eyebrow(sl, no, tag)
    text(sl, M, 86, 1040, 44, title, 30, INK, True)
    text(sl, M, 138, 1040, 22, path, 12, MUTED)

    # 图在哪一侧，文字就换到另一侧
    tx = 700 if side == "left" else M
    ix = M if side == "left" else 620

    # 步骤
    y = 196
    for i, (st, sd) in enumerate(steps):
        rich(sl, tx, y, 500, 22, [("0%d" % (i + 1), 11, DEEP, True), ("　" + st, 14, INK, True)])
        text(sl, tx + 26, y + 26, 470, 60, sd, 11.5, BODY, lh=1.65)
        y += 92
    if note:
        rule(sl, tx, y + 4, 470, 1, HAIR2)
        text(sl, tx, y + 20, 470, 40, "提示　" + note, 11, MUTED, lh=1.6)

    # 图
    iw = 580
    r = ratio_of(img)
    ih = iw / r
    if ih > 450:
        ih = 450
        iw = ih * r
    iy = 170 + (450 - ih) / 2
    picture(sl, img, ix + (580 - iw) / 2, iy, iw)

    footer(sl, label, page, total)
    return sl


def slide_band(prs, no, tag, title, path, steps, img, note, label, page, total):
    sl = blank(prs)
    eyebrow(sl, no, tag)
    text(sl, M, 86, 1040, 44, title, 30, INK, True)
    text(sl, M, 138, 1040, 22, path, 12, MUTED)

    iw = 1120
    r = ratio_of(img)
    ih = iw / r
    if ih > 300:
        ih = 300
        iw = ih * r
    picture(sl, img, M + (CW - iw) / 2, 178, iw)

    sy = 178 + ih + 46
    sy = max(sy, 440)
    colw = (CW - 80) / 3
    for i, (st, sd) in enumerate(steps):
        cx = M + i * (colw + 40)
        rich(sl, cx, sy, colw, 22, [(("0%d" % (i + 1)), 11, DEEP, True), ("　" + st, 13.5, INK, True)])
        text(sl, cx + 24, sy + 25, colw - 24, 60, sd, 11.5, BODY, lh=1.6)
    if note:
        text(sl, M, sy + 108, CW, 22, "提示　" + note, 11, MUTED)

    footer(sl, label, page, total)
    return sl


def slide_close(prs, total):
    sl = blank(prs)
    rule(sl, 0, H - 12, W, 12, DEEP)
    text(sl, M, 150, 900, 50, "需要时，从这一页开始", 32, INK, True)
    text(sl, M, 208, 900, 24, "账号、文档、出问题找谁，都在这。", 15, BODY)

    rows = [
        ("管理员账号", "admin　初始密码见交付说明，登录后请立即修改"),
        ("学员账号", "「系统管理 → 用户管理」里新建，分配角色后即可登录"),
        ("完整说明", "《考试系统操作手册》19 页，含安装部署、排障与默认配置"),
    ]
    y = 300
    for k, v in rows:
        rule(sl, M, y - 18, CW, 1, HAIR2)
        text(sl, M, y, 180, 24, k, 13, MUTED)
        text(sl, M + 200, y - 2, 900, 24, v, 14, INK)
        y += 74
    rule(sl, M, y - 18, CW, 1, HAIR2)

    text(sl, M, 600, CW, 40, "遇到问题先翻手册第 8 章「常见问题与故障排查」；仍解决不了，把页面截图和 server-run.log 一起发过来。", 12, MUTED)
    text(sl, M, 668, 400, 20, "结束", 11, FAINT)
    text(sl, W - M - 120, 668, 120, 20, "%02d / %d" % (total, total), 11, FAINT, align="r")
    return sl


# ---------------------------------------------------------------- 内容

INDEX = [
    ("01", "登录后台", "03"),
    ("02", "首页看什么", "04"),
    ("03", "建大纲", "05"),
    ("04", "建题库、录题", "06"),
    ("05", "组卷策略", "07"),
    ("06", "试卷管理", "08"),
    ("07", "排考试计划", "09"),
    ("08", "报名审核", "10"),
    ("09", "考试监控", "11"),
    ("10", "考试阅卷", "12"),
    ("11", "交卷记录", "13"),
    ("12", "学员报名与答题", "14"),
    ("13", "学员查分", "15"),
]

SLIDES = [
    dict(k="split", no="01", tag="登录", title="登录：一个地址，两种角色",
         path="浏览器访问　http://<服务器IP>:8080",
         steps=[("打开地址", "局域网内用服务器 IP 加 8080 端口访问；部署方式见手册 3.3。"),
                ("填账号密码", "管理员填 admin，学员填管理员分配的登录名；再输入 4 位图形验证码，点验证码可换一张。"),
                ("看菜单认角色", "管理员是完整菜单，学员只有「学员中心」四项，登录后直接落到对应首页。")],
         img="u_login.png", side="right", label="登录",
         note="首次登录后先改密码：「个人中心 → 修改密码」。"),
    dict(k="split", no="02", tag="首页", title="首页：今天的活都在这屏",
         path="登录后默认进入",
         steps=[("看第一行规模数", "学员数、题库量、试题量、试卷数、在考计划、今日考试，一眼看清家底。"),
                ("看第二行运行数", "参考人次、平均分、及格率、待批阅、待审核，卡在哪一环一目了然。"),
                ("点数字进列表", "待办事项在下方逐条列出，点「去处理」直达对应页面。")],
         img="u_dashboard.png", side="left", label="首页",
         note="统计口径来自真实答卷，没交卷就是 0，不做模拟填充。"),
    dict(k="band", no="03", tag="大纲", title="大纲：先立骨架，题才有地方挂",
         path="考试中心 → 大纲管理",
         steps=[("新建大纲", "填名称、分类、版本号，一场考试对应一份大纲。"),
                ("建章节层级", "按章节树往下拆，题目最终挂到最细一级。"),
                ("设为启用", "只有启用状态的大纲，建题库时才能被选到。")],
         img="u_outlines.png", label="大纲管理",
         note="已被题库引用的大纲不能删除，改版请升版本号。"),
    dict(k="split", no="04", tag="题库", title="题库：录一次，出卷只挑不抄",
         path="考试中心 → 题库与试题",
         steps=[("建库并挂大纲", "题库必须挂在大纲下，否则组卷时抽不到题。"),
                ("录题", "六种题型：单选、多选、判断、填空、简答，以及带子题的案例题，同页切换录入。"),
                ("写解析", "解析随成绩展示给学员；主观题另需填参考答案。")],
         img="u_questions.png", side="right", label="题库与试题",
         note="同一道题只存一份，被多张试卷引用时改一次即全部生效。"),
    dict(k="band", no="05", tag="组卷", title="组卷策略：把配比存成模板",
         path="考试中心 → 组卷策略",
         steps=[("设题型配比", "每种题型抽几道、每题多少分，总分自动累加。"),
                ("绑定大纲", "策略从哪份大纲的题库里抽；抽不够会直接提示缺口。"),
                ("保存复用", "同类考试下次套用同一策略，不用重配。")],
         img="u_strategies.png", label="组卷策略",
         note="配比之和要等于试卷总分，否则出卷时报错。"),
    dict(k="band", no="06", tag="试卷", title="试卷：策略一出，卷子就有了",
         path="考试中心 → 试卷管理",
         steps=[("从策略生成", "选策略一键出卷，题号顺序自动打乱。"),
                ("人工调整", "个别题要换就在试卷里手动替换，不动配比。"),
                ("发布试卷", "发布后，才能被考试计划选中。")],
         img="u_papers.png", label="试卷管理",
         note="已被计划引用的试卷删不掉，避免历史答卷断链。"),
    dict(k="band", no="07", tag="排考", title="排考：时间、试卷、谁能考",
         path="考试中心 → 考试计划",
         steps=[("新建计划", "选试卷，填名称、开始与结束时间、考试时长、参考范围。"),
                ("走状态", "草稿 → 已发布 → 进行中 → 已结束，只有「进行中」学员才能进考场。"),
                ("发布成绩", "列表里打开「发布成绩」开关，学员才看得到分数与解析。")],
         img="u_plans.png", label="考试计划",
         note="已发布的计划要改关键信息，建议结束再新建，避免学员端数据错乱。"),
    dict(k="band", no="08", tag="报名", title="报名：要审核就在这里点",
         path="考试中心 → 考试报名",
         steps=[("看申请", "学员在学员端提交报名后，这里出现待审记录。"),
                ("通过或拒绝", "通过后学员状态变「可参加考试」；被拒的尝试开考会被拦住。"),
                ("允许反悔", "学员可自行取消报名，取消后可重新报名。")],
         img="u_signup-review.png", label="考试报名",
         note="计划里已指定参考范围时，名单内的人不需要再走报名审核。"),
    dict(k="band", no="09", tag="监考", title="监考：一个人看一屋子人",
         path="考试中心 → 考试监控",
         steps=[("盯顶部一行", "在考、已交卷、平均分、切屏异常、剩余时间，10 秒自动刷新。"),
                ("找异常", "切屏与窗口失焦按人累计，次数多的考生在列表里标出来。"),
                ("强制交卷", "卡死或确认违规的答卷直接收卷，已答内容照常进阅卷。")],
         img="u_proctor.png", label="考试监控",
         note="切屏和失焦都由浏览器事件采集，学员端不用装任何客户端。"),
    dict(k="band", no="10", tag="阅卷", title="阅卷：主观题逐题给分",
         path="考试中心 → 考试阅卷",
         steps=[("进待批答卷", "列表显示主观题进度，点「批阅」进评分界面。"),
                ("逐题打分", "客观题已自动判分不可改；主观题按参考答案给分并填评语。"),
                ("保存即汇总", "保存后系统重算总分，同步进成绩统计。")],
         img="u_marking.png", label="考试阅卷",
         note="主观题全部批完，答卷才进入「已批阅」，成绩才算定稿。"),
    dict(k="band", no="11", tag="查卷", title="查卷：谁交了、交了什么",
         path="考试中心 → 交卷记录",
         steps=[("三维筛选", "按考试、状态、学员姓名定位一张答卷。"),
                ("展开答卷", "点「答卷」看每题作答与得分，含切屏次数。"),
                ("跳转复核", "需要人工复核的，从抽屉直接跳到阅卷页。")],
         img="u_records.png", label="交卷记录",
         note="记录只增不改，交卷时间以服务端时间为准。"),
    dict(k="band", no="12", tag="学员", title="学员端：报名、答题、交卷",
         path="学员中心 → 考试中心",
         steps=[("报名", "选可参加的考试点「报名」；需审核的等管理员通过。"),
                ("答题", "开始后顶部倒计时以服务端为准，答案自动保存，断网重登可续答。"),
                ("交卷", "交卷前看题号列表的未答标记；交卷幂等，重复点不会出两份卷。")],
         img="u_student.png", label="学员中心",
         note="考试中切屏会被记录，次数多了会进监控端的异常列表。"),
    dict(k="split", no="13", tag="查分", title="查分：总分、对错、解析",
         path="学员中心 → 考试中心 → 成绩",
         steps=[("等发布", "成绩未发布时只显示「已交卷」，属正常现象。"),
                ("看总分", "总分、是否及格、答对题数一屏给出。"),
                ("逐题复盘", "每题列出我的答案、正确答案与解析，错的收进错题本。")],
         img="u_result.png", side="right", label="成绩查询",
         note="错题本按题库归集，可以反复练到会为止。"),
]


def main():
    prs = Presentation()
    prs.slide_width = px(W)
    prs.slide_height = px(H)
    total = 2 + len(SLIDES) + 1

    slide_cover(prs)
    slide_index(prs, INDEX, total)

    page = 3
    for cfg in SLIDES:
        cfg = dict(cfg)
        cfg["page"] = page
        cfg["total"] = total
        if cfg["k"] == "band":
            slide_band(prs, cfg["no"], cfg["tag"], cfg["title"], cfg["path"],
                       cfg["steps"], cfg["img"], cfg["note"], cfg["label"],
                       page, total)
        else:
            slide_split(prs, cfg["no"], cfg["tag"], cfg["title"], cfg["path"],
                        cfg["steps"], cfg["img"], cfg["side"], cfg["note"],
                        cfg["label"], page, total)
        page += 1

    slide_close(prs, total)
    prs.save(OUT)
    print("saved:", OUT, os.path.getsize(OUT), "bytes,", len(prs.slides), "slides")


if __name__ == "__main__":
    main()
